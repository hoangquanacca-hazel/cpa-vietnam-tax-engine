#!/usr/bin/env python3
"""Build the tax/invoice penalty checklist (NĐ 125/2020/NĐ-CP as amended by NĐ 310/2025/NĐ-CP).

Data is hand-transcribed from the source texts in ./sources. Nothing here is generated or
simulated: every row cites article/clause/point and the verify step (verify_catalog.py) fails
if a cited VND range is not present verbatim in the cited source text.
"""
import csv
import json
import pathlib

OUT = pathlib.Path(__file__).parent
M = 1_000_000

OLD_FROM = "2020-12-05"   # NĐ125 hiệu lực
NEW_FROM = "2026-01-16"   # NĐ310 hiệu lực
OLD_TO = "2026-01-15"     # ngày cuối áp dụng phiên bản cũ

TT = "Thủ tục thuế"
KS = "Khai sai / trốn thuế"
HD = "Hóa đơn"

rows = []


def R(lv, dieu, khoan, diem, hv, ht, mn=None, mx=None, dk="", bp="", src="125",
      tu=OLD_FROM, den="", tt="Còn hiệu lực", note="", ty_le=""):
    """ht: CC cảnh cáo | PT phạt tiền theo khung | TL phạt theo tỷ lệ."""
    rows.append(dict(
        lv=lv, dieu=dieu, khoan=khoan, diem=diem, hanh_vi=hv,
        hinh_thuc={"CC": "Cảnh cáo", "PT": "Phạt tiền (khung)", "TL": "Phạt theo tỷ lệ"}[ht],
        min_vnd=None if mn is None else int(mn * M), max_vnd=None if mx is None else int(mx * M),
        ty_le=ty_le, dieu_kien=dk, bien_phap_khac_phuc=bp,
        nguon=src, hieu_luc_tu=tu, hieu_luc_den=den, trang_thai=tt, ghi_chu=note))


REP = "Đã bãi bỏ từ 16/01/2026 (chỉ áp dụng hành vi đã kết thúc trước ngày này)"
SUP = "Đã được thay bằng quy định mới từ 16/01/2026 (chỉ áp dụng hành vi đã kết thúc trước ngày này)"
GMT = "Nghị quyết 107/2023/QH15 (thuế TNDN bổ sung - thuế tối thiểu toàn cầu)"

# ───────────── Điều 10: đăng ký thuế, tạm ngừng ─────────────
R(TT, 10, 1, "", "Đăng ký thuế / thông báo tạm ngừng / thông báo tiếp tục KD trước hạn quá hạn 01-10 ngày", "CC",
  dk="Có tình tiết giảm nhẹ")
R(TT, 10, 1, "b", "Thông báo đơn vị hợp thành chịu trách nhiệm kê khai (GMT) quá hạn", "CC",
  src="310", tu=NEW_FROM, note=GMT)
R(TT, 10, 2, "a", "Đăng ký thuế / thông báo tiếp tục KD trước hạn quá hạn 01-30 ngày", "PT", 1, 2,
  dk="Trừ trường hợp cảnh cáo khoản 1")
R(TT, 10, 2, "b", "Thông báo tạm ngừng hoạt động kinh doanh quá hạn", "PT", 1, 2, dk="Trừ trường hợp cảnh cáo khoản 1")
R(TT, 10, 2, "c", "Không thông báo tạm ngừng hoạt động kinh doanh", "PT", 1, 2)
R(TT, 10, 2, "d", "Không thông báo đơn vị hợp thành chịu trách nhiệm kê khai (GMT)", "PT", 1, 2, src="125+310",
  tu=NEW_FROM, bp="Buộc nộp thông báo đơn vị hợp thành (khoản 5 Điều 10 mới)",
  note="NĐ310 chỉ bổ sung hành vi; khung 1-2 triệu là khung khoản 2 Điều 10 NĐ125. " + GMT)
R(TT, 10, 3, "", "Đăng ký thuế / thông báo tiếp tục KD trước hạn quá hạn 31-90 ngày", "PT", 3, 6)
R(TT, 10, 4, "a", "Đăng ký thuế / thông báo tiếp tục KD trước hạn quá hạn từ 91 ngày", "PT", 6, 10)
R(TT, 10, 4, "b", "Không thông báo tiếp tục KD trước hạn đã thông báo, không phát sinh thuế phải nộp", "PT", 6, 10)

# ───────────── Điều 11: thay đổi thông tin đăng ký thuế ─────────────
R(TT, 11, 1, "a", "Thông báo thay đổi đăng ký thuế quá hạn 01-30 ngày, KHÔNG làm đổi GCN đăng ký thuế/thông báo MST",
  "CC", dk="Có tình tiết giảm nhẹ")
R(TT, 11, 1, "b", "Thông báo thay đổi đăng ký thuế quá hạn 01-10 ngày, CÓ làm đổi GCN đăng ký thuế/thông báo MST",
  "CC", dk="Có tình tiết giảm nhẹ")
R(TT, 11, 2, "", "Thông báo thay đổi đăng ký thuế quá hạn 01-30 ngày, không làm đổi GCN/thông báo MST", "PT", 0.5, 1,
  dk="Trừ trường hợp cảnh cáo điểm a khoản 1")
R(TT, 11, 3, "a", "Thông báo thay đổi quá hạn 31-90 ngày, không làm đổi GCN/thông báo MST", "PT", 1, 3)
R(TT, 11, 3, "b", "Thông báo thay đổi quá hạn 01-30 ngày, có làm đổi GCN/thông báo MST", "PT", 1, 3,
  dk="Trừ trường hợp cảnh cáo điểm b khoản 1")
R(TT, 11, 4, "a", "Thông báo thay đổi quá hạn từ 91 ngày, không làm đổi GCN/thông báo MST", "PT", 3, 5)
R(TT, 11, 4, "b", "Thông báo thay đổi quá hạn 31-90 ngày, có làm đổi GCN/thông báo MST", "PT", 3, 5)
R(TT, 11, 5, "a", "Thông báo thay đổi quá hạn từ 91 ngày, có làm đổi GCN/thông báo MST", "PT", 5, 7)
R(TT, 11, 5, "b", "Không thông báo thay đổi thông tin trong hồ sơ đăng ký thuế", "PT", 5, 7,
  bp="Buộc nộp hồ sơ thay đổi nội dung đăng ký thuế",
  note="Khoản 6 (sửa bởi NĐ310): không xử phạt cá nhân không KD chậm đổi thông tin theo thẻ căn cước, "
       "cơ quan chi trả thu nhập chậm thông báo đổi thông tin căn cước của cá nhân ủy quyền quyết toán, "
       "và thay đổi địa chỉ do thay đổi địa giới hành chính.")

# ───────────── Điều 12: khai sai không dẫn đến thiếu thuế ─────────────
R(KS, 12, 1, "", "Khai sai/thiếu chỉ tiêu trong hồ sơ thuế nhưng không liên quan xác định nghĩa vụ thuế", "PT", 0.5, 1.5,
  dk="Trừ hành vi khoản 2", bp="Buộc khai lại, nộp bổ sung tài liệu")
R(KS, 12, 2, "", "Khai sai/thiếu chỉ tiêu trên tờ khai, phụ lục tờ khai nhưng không liên quan xác định nghĩa vụ thuế",
  "PT", 1.5, 2.5, bp="Buộc khai lại, nộp bổ sung tài liệu")
R(KS, 12, 3, "a", "Khai sai/thiếu chỉ tiêu liên quan xác định nghĩa vụ thuế trong hồ sơ thuế (không thiếu thuế)", "PT", 5, 8,
  bp="Buộc khai lại; điều chỉnh số lỗ, số thuế GTGT đầu vào chuyển kỳ sau")
R(KS, 12, 3, "b", "Khai sai nhưng không thiếu thuế - các hành vi tại khoản 3 Điều 16 và khoản 7 Điều 17", "PT", 5, 8,
  bp="Buộc điều chỉnh số lỗ, số thuế GTGT đầu vào chuyển kỳ sau")

# ───────────── Điều 13: chậm nộp hồ sơ khai thuế ─────────────
R(TT, 13, 1, "", "Nộp hồ sơ khai thuế quá hạn 01-05 ngày", "CC", dk="Có tình tiết giảm nhẹ",
  bp="Buộc nộp tiền chậm nộp nếu chậm nộp hồ sơ dẫn đến chậm nộp thuế")
R(TT, 13, 2, "", "Nộp hồ sơ khai thuế quá hạn 01-30 ngày", "PT", 2, 5, dk="Trừ trường hợp cảnh cáo khoản 1",
  bp="Buộc nộp tiền chậm nộp (nếu chậm nộp thuế)")
R(TT, 13, 3, "", "Nộp hồ sơ khai thuế quá hạn 31-60 ngày", "PT", 5, 8, bp="Buộc nộp tiền chậm nộp (nếu chậm nộp thuế)")
R(TT, 13, 4, "a", "Nộp hồ sơ khai thuế quá hạn 61-90 ngày", "PT", 8, 15, bp="Buộc nộp tiền chậm nộp (nếu chậm nộp thuế)")
R(TT, 13, 4, "b", "Nộp hồ sơ khai thuế quá hạn từ 91 ngày, không phát sinh thuế phải nộp", "PT", 8, 15)
R(TT, 13, 4, "c", "Không nộp hồ sơ khai thuế, không phát sinh thuế phải nộp", "PT", 8, 15, bp="Buộc nộp hồ sơ khai thuế")
R(TT, 13, 4, "d", "Không nộp phụ lục giao dịch liên kết kèm hồ sơ quyết toán TNDN", "PT", 8, 15,
  bp="Buộc nộp phụ lục kèm hồ sơ khai thuế")
C135 = ("Nộp hồ sơ khai thuế quá hạn trên 90 ngày, có phát sinh thuế phải nộp, đã nộp đủ thuế + tiền chậm nộp "
        "trước khi có quyết định kiểm tra/thanh tra hoặc biên bản chậm nộp")
R(TT, 13, 5, "", C135, "PT", 15, 25, src="125", den=OLD_TO, tt=SUP,
  bp="Buộc nộp tiền chậm nộp (khoản 6.a Điều 13 gốc có bao gồm khoản 5)",
  note="Trường hợp nộp trên 90 ngày KHÔNG thuộc khoản 5 bị xử lý như trốn thuế (Điều 17.1.a).")
R(TT, 13, 5, "", C135, "PT", 15, 25, src="310", tu=NEW_FROM,
  bp="Khoản 6.a Điều 13 mới chỉ buộc nộp tiền chậm nộp đối với khoản 1, 2, 3, 4 (không còn khoản 5)",
  note="Từ 16/01/2026: nếu tiền phạt > số thuế phát sinh trên hồ sơ (hoặc tổng thuế các hồ sơ nộp cùng ngày cùng "
       "sắc thuế) thì mức phạt tối đa = số thuế phát sinh, nhưng không thấp hơn mức trung bình khung khoản 4 (11,5 triệu). "
       "Trigger đổi thành 'cơ quan khác công bố quyết định thanh tra, kiểm tra'.")

# ───────────── Điều 14: cung cấp thông tin ─────────────
R(TT, 14, 1, "a", "Cung cấp thông tin, hồ sơ pháp lý đăng ký thuế theo thông báo của cơ quan thuế quá hạn từ 05 ngày làm việc",
  "PT", 2, 3)
R(TT, 14, 1, "b", "Cung cấp thông tin, tài liệu, sổ kế toán xác định nghĩa vụ thuế theo thông báo quá hạn từ 05 ngày làm việc",
  "PT", 2, 3)
R(TT, 14, 2, "a", "Không cung cấp / cung cấp không đầy đủ, không chính xác thông tin, chứng từ, hóa đơn, sổ kế toán xác định "
  "nghĩa vụ thuế; số hiệu, số dư tài khoản", "PT", 3, 5, bp="Buộc cung cấp thông tin")
R(TT, 14, 2, "b", "Không cung cấp / cung cấp không đầy đủ, không đúng chỉ tiêu nghĩa vụ thuế phải đăng ký, không làm giảm "
  "nghĩa vụ thuế", "PT", 3, 5, bp="Buộc cung cấp thông tin")
R(TT, 14, 2, "c", "Không cung cấp / cung cấp không đầy đủ thông tin tài khoản tại TCTD, KBNN, công nợ bên thứ ba khi cơ quan "
  "thuế yêu cầu", "PT", 3, 5, bp="Buộc cung cấp thông tin")

# ───────────── Điều 15: kiểm tra, thanh tra, cưỡng chế ─────────────
for d, t in [("a", "Không nhận quyết định thanh tra, kiểm tra thuế, quyết định cưỡng chế khi cơ quan thuế giao, gửi"),
             ("b", "Không chấp hành quyết định thanh tra, kiểm tra thuế quá thời hạn 03 ngày làm việc trở lên"),
             ("c", "Cung cấp hồ sơ, tài liệu, hóa đơn, sổ kế toán quá 06 giờ làm việc khi kiểm tra tại trụ sở"),
             ("d", "Cung cấp không đầy đủ, chính xác thông tin, tài liệu, sổ kế toán khi kiểm tra tại trụ sở"),
             ("đ", "Không ký biên bản kiểm tra, thanh tra thuế trong 05 ngày làm việc")]:
    R(TT, 15, 1, d, t, "PT", 2, 5, bp="Buộc cung cấp thông tin, tài liệu, sổ kế toán" if d == "d" else "")
R(TT, 15, 2, "a", "Không cung cấp số liệu, tài liệu, sổ kế toán khi cơ quan có thẩm quyền yêu cầu trong kiểm tra/thanh tra tại trụ sở",
  "PT", 5, 10, bp="Buộc cung cấp thông tin, tài liệu, sổ kế toán")
R(TT, 15, 2, "b", "Không thực hiện / thực hiện không đúng quyết định niêm phong hồ sơ, két quỹ, kho hàng...", "PT", 5, 10)
R(TT, 15, 2, "c", "Tự ý tháo bỏ, thay đổi dấu hiệu niêm phong hợp pháp", "PT", 5, 10)

# ───────────── Điều 16: khai sai dẫn đến thiếu thuế (20%) ─────────────
T20 = "20% số thuế khai thiếu / số thuế miễn, giảm, hoàn cao hơn quy định"
BP16 = "Buộc nộp đủ thuế thiếu + tiền chậm nộp; điều chỉnh lỗ, thuế GTGT đầu vào chuyển kỳ sau"
R(KS, 16, 1, "a", "Khai sai căn cứ tính thuế / số thuế được khấu trừ / sai trường hợp miễn giảm hoàn - nghiệp vụ đã phản ánh "
  "đầy đủ trong sổ sách, hóa đơn, chứng từ hợp pháp", "TL", ty_le=T20, bp=BP16)
R(KS, 16, 1, "b", "Khai sai làm giảm thuế phải nộp (không thuộc điểm a), đã tự khai bổ sung và nộp đủ thuế trước khi cơ quan "
  "thuế kết thúc thanh tra, kiểm tra tại trụ sở", "TL", ty_le=T20, bp=BP16, tt=REP, den=OLD_TO,
  note="Điểm này bị NĐ310 bãi bỏ; thay bằng Điều 9.3: khai bổ sung + nộp đủ thuế TRƯỚC khi có quyết định "
       "kiểm tra/thanh tra hoặc trước khi cơ quan thuế phát hiện thì KHÔNG bị xử phạt.")
C161 = ("Khai sai bị xác định là trốn thuế nhưng vi phạm lần đầu về trốn thuế, đã khai bổ sung và nộp đủ thuế trước "
        "khi có quyết định xử phạt, cơ quan thuế đã lập biên bản ghi nhận là khai sai")
R(KS, 16, 1, "c", C161 + " (biên bản thanh tra, kiểm tra thuế hoặc biên bản vi phạm hành chính)", "TL", ty_le=T20, bp=BP16,
  den=OLD_TO, tt=SUP)
R(KS, 16, 1, "c", C161 + " (chỉ còn: biên bản vi phạm hành chính)", "TL", ty_le=T20, bp=BP16, src="310", tu=NEW_FROM,
  note="Bản OCR NĐ310 ghi tiêu đề 'điểm e'; nội dung trích khớp điểm c khoản 1 Điều 16 gốc (khoản 1 không có điểm e).")
R(KS, 16, 1, "d", "Khai sai dẫn đến thiếu thuế với giao dịch liên kết nhưng đã lập hồ sơ xác định giá thị trường / gửi phụ lục "
  "giao dịch liên kết", "TL", ty_le=T20, bp=BP16)
R(KS, 16, 1, "đ", "Sử dụng hóa đơn, chứng từ không hợp pháp để hạch toán đầu vào, nhưng người mua chứng minh được lỗi thuộc bên "
  "bán và đã hạch toán đầy đủ", "TL", ty_le=T20, bp=BP16)

# ───────────── Điều 17: trốn thuế ─────────────
BP17 = "Buộc nộp đủ thuế trốn (+ tiền chậm nộp); điều chỉnh lỗ, thuế GTGT đầu vào"
TRON = "Phạt 1-3 lần số thuế trốn (mức cụ thể theo tình tiết - xem sheet Nguyên tắc chung)"
for d, t in [("a", "Không nộp hồ sơ đăng ký thuế; không nộp hồ sơ khai thuế hoặc nộp sau 90 ngày kể từ ngày hết hạn"),
             ("b", "Không ghi sổ kế toán các khoản thu; không khai, khai sai dẫn đến thiếu thuế (trừ Điều 16)"),
             ("c", "Không lập hóa đơn khi bán hàng hóa, dịch vụ (trừ trường hợp đã khai thuế); lập hóa đơn sai số lượng, giá trị "
                   "để khai thuế thấp hơn thực tế và bị phát hiện sau hạn nộp hồ sơ khai thuế"),
             ("d", "Sử dụng hóa đơn không hợp pháp; sử dụng không hợp pháp hóa đơn để khai thuế làm giảm thuế phải nộp / tăng thuế "
                   "được hoàn, miễn, giảm"),
             ("đ", "Sử dụng chứng từ không hợp pháp; chứng từ, tài liệu không phản ánh đúng bản chất/giá trị giao dịch; hủy vật tư, "
                   "hàng hóa không đúng thực tế để giảm thuế"),
             ("e", "Sử dụng hàng hóa không chịu thuế/miễn thuế sai mục đích mà không khai báo chuyển đổi mục đích"),
             ("g", "Có hoạt động KD trong thời gian xin ngừng/tạm ngừng mà không thông báo cơ quan thuế")]:
    R(KS, 17, 1, d, t, "TL", ty_le=TRON, bp=BP17,
      dk={"a": "Trừ trường hợp điểm b, c khoản 4 và khoản 5 Điều 13",
          "g": "Trừ trường hợp điểm b khoản 4 Điều 10"}.get(d, ""),
      note=("Nếu bị phát hiện sau hạn nộp hồ sơ khai thuế nhưng KHÔNG làm giảm thuế phải nộp: xử phạt theo Điều 12.3 "
            "(5-8 triệu) - khoản 7 Điều 17." if d in ("b", "đ", "e") else ""))

# ───────────── Điều 18: ngân hàng, người bảo lãnh ─────────────
R(KS, 18, 1, "", "Ngân hàng thương mại không trích chuyển tiền từ tài khoản người nộp thuế vào NSNN theo yêu cầu của cơ quan thuế",
  "TL", ty_le="Phạt bằng số thuế + tiền chậm nộp + tiền phạt không trích chuyển (trừ số dư tối thiểu)",
  dk="Trừ trường hợp tài khoản không còn số dư hoặc đã trích chuyển toàn bộ số dư mà vẫn không đủ số tiền phải nộp",
  note="Đối tượng: ngân hàng thương mại, không phải người nộp thuế.")

# ───────────── Điều 19: tổ chức, cá nhân liên quan ─────────────
R(TT, 19, 1, "", "Cung cấp thông tin, tài liệu xác định nghĩa vụ thuế / tài khoản của người nộp thuế theo yêu cầu cơ quan thuế quá "
  "hạn từ 05 ngày", "PT", 2, 6, tu=OLD_FROM, note="Giữ nguyên khung ở bản cũ và bản mới.", src="125+310")
R(TT, 19, 2, "a", "Thông đồng, bao che người nộp thuế trốn thuế; không thực hiện quyết định cưỡng chế hành chính thuế "
  "(trừ Điều 18)", "PT", 6, 16, den=OLD_TO, tt=SUP)
R(TT, 19, 2, "b", "Không cung cấp / cung cấp không chính xác thông tin tài sản, quyền, nghĩa vụ tài sản của người nộp thuế; tài "
  "khoản tại TCTD, KBNN", "PT", 6, 16, den=OLD_TO, tt=SUP, bp="Buộc cung cấp thông tin")
for d, t in [("a", "Cung cấp KHÔNG CHÍNH XÁC thông tin tài sản, quyền, nghĩa vụ tài sản của người nộp thuế hoặc thông tin xác định "
                   "nghĩa vụ thuế"),
             ("b", "Cung cấp KHÔNG CHÍNH XÁC thông tin tài khoản của người nộp thuế tại TCTD, KBNN, chi nhánh ngân hàng nước ngoài"),
             ("c", "Cung cấp KHÔNG CHÍNH XÁC thông tin tiền lương, tiền công, thu nhập của người nộp thuế")]:
    R(TT, 19, 2, d, t, "PT", 6, 10, src="310", tu=NEW_FROM, bp="Buộc cung cấp thông tin đầy đủ, chính xác")
for d, t in [("a", "KHÔNG cung cấp thông tin tài sản, quyền, nghĩa vụ tài sản của người nộp thuế hoặc thông tin xác định nghĩa vụ thuế"),
             ("b", "KHÔNG cung cấp thông tin tài khoản của người nộp thuế tại TCTD, KBNN, chi nhánh ngân hàng nước ngoài"),
             ("c", "KHÔNG cung cấp thông tin tiền lương, tiền công, thu nhập của người nộp thuế")]:
    R(TT, 19, 3, d, t, "PT", 10, 16, src="310", tu=NEW_FROM, bp="Buộc cung cấp thông tin đầy đủ, chính xác")
R(TT, 19, 3, "đ", "Thông đồng, bao che người nộp thuế trốn thuế; không thực hiện quyết định cưỡng chế hành chính thuế (trừ Điều 18)",
  "PT", 10, 16, src="310", tu=NEW_FROM)

# ───────────── Chương III: HÓA ĐƠN ─────────────
# Điều 20, 21, 23: bãi bỏ hoàn toàn
for d, k, dm, t, a, b in [
    (20, 1, "", "Không ký hợp đồng in bằng văn bản / tổ chức in in hóa đơn đặt in không có quyết định in", 0.5, 1.5),
    (20, 2, "", "Đặt in hóa đơn khi cơ quan thuế đã có văn bản thông báo không đủ điều kiện đặt in (trừ trường hợp cơ quan thuế không có ý kiến bằng văn bản khi nhận đề nghị sử dụng hóa đơn đặt in)", 2, 4),
    (20, 3, "", "Đặt in hóa đơn theo mẫu đã phát hành của tổ chức/cá nhân khác hoặc đặt in trùng số cùng ký hiệu", 20, 50),
    (21, 2, "", "In hóa đơn đặt in mà không ký hợp đồng in bằng văn bản", 0.5, 1.5),
    (21, 3, "", "Báo cáo về việc in hóa đơn quá hạn từ 06 ngày trở lên (trừ trường hợp cảnh cáo điểm b khoản 1)", 2, 4),
    (21, 4, "", "Không hủy sản phẩm in hỏng, in thừa khi thanh lý hợp đồng in", 4, 8),
    (21, 5, "a", "Nhận in hóa đơn đặt in khi không đủ điều kiện in hóa đơn", 6, 18),
    (21, 5, "b", "Không khai báo việc làm mất hóa đơn trước khi giao cho khách hàng", 6, 18),
    (21, 6, "", "Chuyển nhượng toàn bộ hoặc một khâu trong hợp đồng in hóa đơn cho cơ sở in khác", 10, 20),
    (21, 7, "", "In hóa đơn theo mẫu đã phát hành của tổ chức/cá nhân khác hoặc in trùng số cùng ký hiệu", 20, 50),
    (22, 1, "a", "Cho, bán hóa đơn đặt in chưa phát hành", 15, 45),
    (22, 1, "b", "Cho, bán hóa đơn đặt in của khách hàng đặt in hóa đơn cho tổ chức, cá nhân khác", 15, 45),
    (23, 1, "a", "Nộp thông báo điều chỉnh thông tin phát hành hóa đơn (đổi địa chỉ/tên) quá hạn 10-20 ngày", 0.5, 1.5),
    (23, 1, "b", "Nộp bảng kê hóa đơn chưa sử dụng khi đổi địa chỉ quá hạn 10-20 ngày", 0.5, 1.5),
    (23, 1, "c", "Sử dụng hóa đơn đã thông báo phát hành nhưng chưa đến hạn sử dụng", 0.5, 1.5),
    (23, 2, "a", "Lập thông báo phát hành không đầy đủ nội dung, đã bị cơ quan thuế thông báo điều chỉnh nhưng vẫn lập hóa đơn", 2, 4),
    (23, 2, "b", "Không niêm yết thông báo phát hành hóa đơn đúng quy định", 2, 4),
    (23, 2, "c", "Nộp thông báo điều chỉnh thông tin phát hành hóa đơn (đổi địa chỉ/tên) quá hạn từ 21 ngày", 2, 4),
    (23, 2, "d", "Nộp bảng kê hóa đơn chưa sử dụng khi đổi địa chỉ quá hạn từ 21 ngày", 2, 4),
    (23, 3, "", "Không lập thông báo phát hành hóa đơn trước khi đưa vào sử dụng (hóa đơn gắn nghiệp vụ đã khai, nộp thuế hoặc chưa đến kỳ kê khai, nộp thuế)", 6, 18)]:
    R(HD, d, k, dm, t, "PT", a, b, tt=REP, den=OLD_TO,
      note=f"Điều {d} NĐ125 bị bãi bỏ hoàn toàn bởi Điều 2 NĐ310." if d != 22 else "Khoản 1 Điều 22 bị bãi bỏ bởi NĐ310.")
R(HD, 21, 1, "", "Báo cáo về việc nhận in hóa đơn quá hạn 01-05 ngày (và 06-10 ngày nếu có tình tiết giảm nhẹ)", "CC",
  tt=REP, den=OLD_TO, note="Điều 21 NĐ125 bị bãi bỏ bởi Điều 2 NĐ310.")

# Điều 22: cho, bán hóa đơn
R(HD, 22, 2, "", "Cho, bán hóa đơn mua của cơ quan thuế nhưng chưa lập", "PT", 20, 50, den=OLD_TO, tt=SUP,
  bp="Buộc hủy hóa đơn; buộc nộp lại số lợi bất hợp pháp")
R(HD, 22, 2, "", "Cho, bán hóa đơn", "PT", 20, 50, src="310", tu=NEW_FROM,
  bp="Buộc nộp lại số lợi bất hợp pháp (điểm b khoản 3 Điều 22 NĐ125)")

# Điều 24: lập hóa đơn
R(HD, 24, 1, "a", "Lập hóa đơn không đúng thời điểm nhưng không dẫn đến chậm nghĩa vụ thuế", "CC", den=OLD_TO, tt=REP,
  dk="Có tình tiết giảm nhẹ")
R(HD, 24, 1, "b", "Lập hóa đơn liên tục từ số nhỏ đến lớn nhưng khác quyển; đã hủy quyển số nhỏ sau khi phát hiện", "CC")
R(HD, 24, 1, "c", "Lập sai loại hóa đơn nhưng đã lập lại đúng trước khi có quyết định thanh tra/kiểm tra, không ảnh hưởng nghĩa "
  "vụ thuế", "CC")
R(HD, 24, 2, "a", "Không lập hóa đơn tổng hợp theo quy định", "PT", 0.5, 1.5, den=OLD_TO, tt=SUP,
  note="Khoản 2 cũ bị thay hoàn toàn bởi NĐ310; hành vi 'không lập hóa đơn tổng hợp' không còn nêu riêng.")
R(HD, 24, 2, "b", "Không lập hóa đơn cho hàng khuyến mại, quảng cáo, hàng mẫu, cho biếu tặng, trả lương", "PT", 0.5, 1.5,
  den=OLD_TO, tt=SUP)
R(HD, 24, 3, "", "Lập hóa đơn không đúng thời điểm nhưng không dẫn đến chậm nghĩa vụ thuế", "PT", 3, 5, den=OLD_TO, tt=SUP)
R(HD, 24, 4, "a", "Lập hóa đơn không đúng thời điểm theo quy định (trừ điểm a khoản 1, khoản 3)", "PT", 4, 8, den=OLD_TO, tt=REP)
R(HD, 24, 4, "b", "Lập hóa đơn không theo thứ tự từ số nhỏ đến số lớn", "PT", 4, 8, dk="Trừ trường hợp cảnh cáo điểm b khoản 1")
R(HD, 24, 4, "c", "Lập hóa đơn ghi ngày trước ngày mua hóa đơn của cơ quan thuế", "PT", 4, 8)
R(HD, 24, 4, "d", "Lập sai loại hóa đơn đã giao cho người mua hoặc đã kê khai thuế", "PT", 4, 8,
  dk="Trừ trường hợp cảnh cáo điểm c khoản 1",
  bp="Buộc lập hóa đơn theo quy định (đến 15/01/2026: khi người mua yêu cầu; từ 16/01/2026 không còn điều kiện này)")
R(HD, 24, 4, "đ", "Lập hóa đơn điện tử khi chưa có thông báo chấp thuận của cơ quan thuế", "PT", 4, 8)
R(HD, 24, 4, "e", "Lập hóa đơn trong thời gian tạm ngừng hoạt động kinh doanh (trừ hợp đồng ký trước)", "PT", 4, 8)
R(HD, 24, 4, "g", "Lập hóa đơn điện tử từ máy tính tiền không kết nối, chuyển dữ liệu điện tử với cơ quan thuế", "PT", 4, 8)
R(HD, 24, 5, "", "Không lập hóa đơn khi bán hàng hóa, cung cấp dịch vụ cho người mua theo quy định", "PT", 10, 20, den=OLD_TO,
  tt=REP, bp="Buộc lập hóa đơn khi người mua yêu cầu",
  note="Khoản 5 bị bãi bỏ bởi NĐ310; hành vi không lập hóa đơn nay thuộc khoản 3 mới (theo số lượng hóa đơn).")

NHOM_A = "Nhóm A: hàng KM/quảng cáo/mẫu, cho biếu tặng, trả lương, tiêu dùng nội bộ, cho vay mượn/hoàn trả"
NHOM_B = "Nhóm B: bán hàng hóa, cung cấp dịch vụ"
# Mới - Điều 24.2: lập hóa đơn KHÔNG ĐÚNG THỜI ĐIỂM (theo số lượng hóa đơn)
tiers_late = [
    ("a", None, None, f"{NHOM_A}: 01 số hóa đơn", "CC"),
    ("b", 0.5, 1.5, f"{NHOM_A}: 02 đến dưới 10 số; {NHOM_B}: 01 số", "PT"),
    ("c", 2, 5, f"{NHOM_A}: 10 đến dưới 50 số; {NHOM_B}: 02 đến dưới 10 số", "PT"),
    ("d", 5, 15, f"{NHOM_A}: 50 đến dưới 100 số; {NHOM_B}: 10 đến dưới 20 số", "PT"),
    ("đ", 15, 30, f"{NHOM_A}: từ 100 số; {NHOM_B}: 20 đến dưới 50 số", "PT"),
    ("e", 30, 50, f"{NHOM_B}: 50 đến dưới 100 số", "PT"),
    ("g", 50, 70, f"{NHOM_B}: từ 100 số trở lên", "PT")]
for d, a, b, dk, ht in tiers_late:
    R(HD, 24, 2, d, "Lập hóa đơn KHÔNG ĐÚNG THỜI ĐIỂM theo quy định", ht, a, b, dk=dk, src="310", tu=NEW_FROM,
      note="Nhiều hành vi trong cùng một vụ việc chỉ bị xử phạt một hành vi theo khung tương ứng tổng số hóa đơn "
           "(điểm đ khoản 3 Điều 5 NĐ125 mới).")
tiers_none = [
    ("a", None, None, f"{NHOM_A}: 01 số hóa đơn", "CC"),
    ("b", 1, 2, f"{NHOM_A}: 02 đến dưới 10 số; {NHOM_B}: 01 số", "PT"),
    ("c", 2, 10, f"{NHOM_A}: 10 đến dưới 50 số; {NHOM_B}: 02 đến dưới 10 số", "PT"),
    ("d", 10, 30, f"{NHOM_A}: 50 đến dưới 100 số; {NHOM_B}: 10 đến dưới 20 số", "PT"),
    ("đ", 30, 50, f"{NHOM_A}: từ 100 số; {NHOM_B}: 20 đến dưới 50 số", "PT"),
    ("e", 60, 80, f"{NHOM_B}: từ 50 số hóa đơn trở lên", "PT")]
for d, a, b, dk, ht in tiers_none:
    R(HD, 24, 3, d, "KHÔNG LẬP hóa đơn theo quy định", ht, a, b, dk=dk, src="310", tu=NEW_FROM,
      bp="Buộc lập hóa đơn theo quy định",
      note="Trường hợp không lập hóa đơn để trốn thuế: xem Điều 17.1.c (phạt 1-3 lần số thuế trốn). "
           "Điểm e khoản 3 Điều 5 NĐ125 (sửa bởi NĐ310): nhiều hành vi không lập hóa đơn trong một vụ việc chỉ phạt một hành vi.")

# Điều 25
R(HD, 25, 1, "", "Khai báo mất, cháy, hỏng hóa đơn quá hạn 01-05 ngày", "CC", dk="Có tình tiết giảm nhẹ", src="125+310",
  note="NĐ310 chỉ sửa tên Điều 25.")
R(HD, 25, 2, "", "Khai báo mất, cháy, hỏng hóa đơn quá hạn 01-05 ngày", "PT", 1, 4, dk="Trừ trường hợp cảnh cáo khoản 1")
R(HD, 25, 3, "a", "Khai báo mất, cháy, hỏng hóa đơn quá hạn từ 06 ngày", "PT", 4, 8)
R(HD, 25, 3, "b", "Không khai báo mất, cháy, hỏng hóa đơn", "PT", 4, 8)

# Điều 26
R(HD, 26, 1, "a", "Làm mất, cháy, hỏng hóa đơn đã lập (trừ liên giao khách), đã kê khai thuế, có chứng từ chứng minh", "CC",
  dk="Có tình tiết giảm nhẹ")
R(HD, 26, 1, "b", "Làm mất, cháy, hỏng hóa đơn đã lập sai, đã xóa bỏ và đã lập hóa đơn thay thế", "CC")
R(HD, 26, 2, "", "Làm mất, cháy, hỏng liên giao khách hàng của hóa đơn đã lập; người bán đã kê khai, có chứng từ", "PT", 3, 5,
  dk="Có tình tiết giảm nhẹ")
R(HD, 26, 3, "a", "Làm mất, cháy, hỏng hóa đơn đã phát hành, đã mua của cơ quan thuế nhưng chưa lập", "PT", 4, 8,
  den=OLD_TO, tt=SUP)
R(HD, 26, 3, "a", "Làm mất, cháy, hỏng hóa đơn mua của cơ quan thuế nhưng chưa lập", "PT", 4, 8, src="310", tu=NEW_FROM,
  note="NĐ310 bỏ cụm 'đã phát hành' - phạm vi hẹp hơn bản cũ.")
R(HD, 26, 3, "b", "Làm mất, cháy, hỏng liên giao khách hàng của hóa đơn đã lập; người bán đã kê khai, có chứng từ", "PT", 4, 8)
R(HD, 26, 4, "", "Làm mất, cháy, hỏng hóa đơn đã lập, đã khai, nộp thuế (trừ các trường hợp khoản 1, 2, 3)", "PT", 5, 10)

# Điều 27
R(HD, 27, 1, "", "Hủy/tiêu hủy hóa đơn quá hạn 01-05 ngày làm việc", "CC", dk="Có tình tiết giảm nhẹ", den=OLD_TO, tt=SUP)
R(HD, 27, 1, "", "Tiêu hủy hóa đơn quá hạn 01-05 ngày làm việc", "CC", dk="Có tình tiết giảm nhẹ", src="310", tu=NEW_FROM)
for d, t, tt_, den_ in [
        ("a", "Hủy không đúng quy định hóa đơn đã phát hành nhưng chưa lập, hóa đơn không còn giá trị sử dụng", SUP, OLD_TO),
        ("b", "Không hủy hóa đơn đã phát hành nhưng chưa lập, không còn giá trị sử dụng; không hủy hóa đơn mua của cơ quan thuế "
              "đã hết hạn sử dụng", SUP, OLD_TO)]:
    R(HD, 27, 2, d, t, "PT", 2, 4, tt=tt_, den=den_, bp="Buộc hủy hóa đơn" if d == "b" else "")
R(HD, 27, 2, "a", "Tiêu hủy không đúng quy định hóa đơn đặt in mua của cơ quan thuế không tiếp tục sử dụng, không còn giá trị",
  "PT", 2, 4, src="310", tu=NEW_FROM)
R(HD, 27, 2, "b", "Không tiêu hủy hóa đơn đặt in mua của cơ quan thuế không tiếp tục sử dụng, không còn giá trị", "PT", 2, 4,
  src="310", tu=NEW_FROM, bp="Buộc tiêu hủy hóa đơn")
R(HD, 27, 2, "c", "Hủy/tiêu hủy hóa đơn quá hạn 01-10 ngày làm việc", "PT", 2, 4, dk="Trừ trường hợp cảnh cáo khoản 1",
  den=OLD_TO, tt=SUP)
R(HD, 27, 2, "c", "Tiêu hủy hóa đơn quá hạn 01-10 ngày làm việc", "PT", 2, 4, dk="Trừ trường hợp cảnh cáo khoản 1",
  src="310", tu=NEW_FROM)
for d, t in [("a", "Hủy/tiêu hủy hóa đơn quá hạn từ 11 ngày làm việc"),
             ("b", "Không hủy, không tiêu hủy hóa đơn theo quy định"),
             ("đ", "Hủy/tiêu hủy hóa đơn không đúng trình tự, thủ tục"),
             ("e", "Tiêu hủy hóa đơn không đúng các trường hợp phải tiêu hủy")]:
    R(HD, 27, 3, d, t, "PT", 4, 8, tt=SUP, den=OLD_TO, bp="Buộc hủy, tiêu hủy hóa đơn" if d == "b" else "")
R(HD, 27, 3, "c", "Không hủy hóa đơn điện tử lập sai sót sau khi quá hạn cơ quan thuế thông báo kiểm tra sai, sót", "PT", 4, 8,
  tt=REP, den=OLD_TO, bp="Buộc hủy hóa đơn")
R(HD, 27, 3, "d", "Không hủy hóa đơn đặt in chưa phát hành nhưng không còn sử dụng", "PT", 4, 8, tt=REP, den=OLD_TO,
  bp="Buộc hủy hóa đơn")
for d, t in [("a", "Tiêu hủy hóa đơn quá hạn từ 11 ngày làm việc"),
             ("b", "Không tiêu hủy hóa đơn theo quy định"),
             ("c", "Tiêu hủy hóa đơn không đúng trình tự, thủ tục"),
             ("đ", "Tiêu hủy hóa đơn không đúng các trường hợp phải tiêu hủy")]:
    R(HD, 27, 3, d, t, "PT", 4, 8, src="310", tu=NEW_FROM, bp="Buộc tiêu hủy hóa đơn" if d == "b" else "",
      note="Ký hiệu điểm theo bản OCR NĐ310 (đánh số lại a-d, đ); cần đối chiếu với bản gốc.")

# Điều 28
R(HD, 28, 1, "", "Sử dụng hóa đơn không hợp pháp / sử dụng không hợp pháp hóa đơn (Điều 4 NĐ125)", "PT", 20, 50,
  dk="Trừ điểm đ khoản 1 Điều 16 và điểm d khoản 1 Điều 17 (khi thuộc diện phạt 20% hoặc trốn thuế)",
  bp="Đến 15/01/2026: buộc hủy hóa đơn đã sử dụng (khoản 2 Điều 28 bị bãi bỏ từ 16/01/2026)")

# Điều 29
R(HD, 29, 1, "", "Nộp thông báo, báo cáo về hóa đơn quá hạn 01-05 ngày", "CC", dk="Có tình tiết giảm nhẹ")
R(HD, 29, 2, "a", "Nộp thông báo, báo cáo về hóa đơn quá hạn 01-10 ngày", "PT", 1, 3, dk="Trừ trường hợp cảnh cáo khoản 1")
R(HD, 29, 2, "b", "Lập sai / không đầy đủ nội dung thông báo, báo cáo về hóa đơn", "PT", 1, 3,
  bp="Buộc lập, gửi thông báo, báo cáo về hóa đơn",
  note="Không phạt nếu tự phát hiện, lập lại báo cáo thay thế trước khi có quyết định thanh tra/kiểm tra.")
R(HD, 29, 3, "", "Nộp thông báo, báo cáo về hóa đơn quá hạn 11-20 ngày", "PT", 2, 4)
R(HD, 29, 4, "", "Nộp thông báo, báo cáo về hóa đơn quá hạn 21-90 ngày", "PT", 4, 8)
R(HD, 29, 5, "a", "Nộp thông báo, báo cáo về hóa đơn quá hạn từ 91 ngày", "PT", 5, 15)
R(HD, 29, 5, "b", "Không nộp thông báo, báo cáo về hóa đơn", "PT", 5, 15, bp="Buộc lập, gửi thông báo, báo cáo về hóa đơn")

# Điều 30
R(HD, 30, 1, "", "Chuyển dữ liệu hóa đơn điện tử cho cơ quan thuế quá hạn 01-05 ngày làm việc", "PT", 2, 5)
R(HD, 30, 2, "a", "Chuyển dữ liệu hóa đơn điện tử quá hạn 06-10 ngày làm việc", "PT", 5, 8)
R(HD, 30, 2, "b", "Chuyển bảng tổng hợp dữ liệu hóa đơn điện tử không đầy đủ số lượng hóa đơn đã lập trong kỳ", "PT", 5, 8,
  bp="Buộc chuyển dữ liệu hóa đơn điện tử")
R(HD, 30, 3, "a", "Chuyển dữ liệu hóa đơn điện tử quá hạn từ 11 ngày làm việc", "PT", 10, 20)
R(HD, 30, 3, "b", "Không chuyển dữ liệu hóa đơn điện tử cho cơ quan thuế theo thời hạn", "PT", 10, 20,
  bp="Buộc chuyển dữ liệu hóa đơn điện tử")

# Điều 31
R(HD, 31, 1, "", "Cung cấp phần mềm hóa đơn tự in không đảm bảo nguyên tắc / in ra không đủ nội dung", "PT", 4, 8,
  den=OLD_TO, tt=SUP, note="Dành cho tổ chức cung cấp dịch vụ/phần mềm hóa đơn.")
R(HD, 31, 2, "", "Cung cấp phần mềm hóa đơn điện tử không đảm bảo nguyên tắc", "PT", 4, 8, den=OLD_TO, tt=SUP,
  note="Dành cho tổ chức cung cấp dịch vụ/phần mềm hóa đơn.")
R(HD, 31, "", "", "Cung cấp giải pháp khởi tạo, kết nối, nhận, truyền, lưu trữ, xử lý dữ liệu hóa đơn điện tử không đảm bảo "
  "nguyên tắc", "PT", 4, 8, src="310+125", tu=NEW_FROM,
  note="NĐ310 thay Điều 31. Số tiền trong file OCR NĐ310 bị lỗi ('§.000.000'); khung 4-8 triệu được đối chiếu với "
       "Điều 31 NĐ125 gốc (cùng khung) - cần xác nhận lại trên bản PDF NĐ310.")

# ───────────── Nguyên tắc chung ─────────────
general = [
    ("Mức tối đa mỗi hành vi (đến hết 15/01/2026)", "Thủ tục thuế - tổ chức: tối đa 200.000.000 đ; cá nhân/hộ KD: tối đa "
     "100.000.000 đ. Hóa đơn - tổ chức: tối đa 100.000.000 đ; cá nhân/hộ KD: tối đa 50.000.000 đ. Hình thức bổ sung: đình chỉ "
     "hoạt động in hóa đơn (khoản 2 Điều 7).", "Điều 7.1.b, 7.2, 7.4.a NĐ125 gốc"),
    ("Mức phạt tối đa mỗi hành vi (từ 16/01/2026)", "NĐ310 thay điểm b khoản 1 Điều 7: mức phạt tiền tối đa đối với hành vi vi "
     "phạm thủ tục thuế và hóa đơn thực hiện theo pháp luật về xử lý vi phạm hành chính (con số cụ thể KHÔNG nêu trong NĐ310 và "
     "không có trong nguồn đã cung cấp). Khoản 2 Điều 7 (đình chỉ in hóa đơn) bị bãi bỏ.", "Điều 1.5.a, Điều 2.2 NĐ310"),
    ("Mức phạt cá nhân so với tổ chức", "Khung phạt trong bảng là mức đối với TỔ CHỨC. Với cùng một hành vi, mức phạt tổ chức gấp "
     "02 lần mức phạt cá nhân, TRỪ hành vi tại Điều 16, 17, 18. Hộ gia đình, hộ kinh doanh áp dụng như cá nhân. Cột 'cá nhân' "
     "trong Excel = 1/2 mức tổ chức; không tính cho hành vi Điều 16, 17, 18 (phạt theo tỷ lệ). Điều 19 sau sửa đổi NĐ310 không "
     "còn khoản nêu rõ nguyên tắc này - nên đối chiếu thêm.", "Điều 5.5, Điều 7.4.a NĐ125 (NĐ310 không sửa Điều 5.5)"),
    ("Mức phạt cụ thể trong khung (từ 16/01/2026)", "Mức phạt cụ thể = mức TRUNG BÌNH của khung. 01 tình tiết giảm nhẹ: giảm "
     "10% mức trung bình; 01 tình tiết tăng nặng: tăng 10%; từ 02 tình tiết giảm nhẹ trở lên: áp mức TỐI THIỂU của khung; "
     "từ 02 tình tiết tăng nặng trở lên: áp mức TỐI ĐA của khung.", "Điều 7.4.d NĐ125 (sửa bởi NĐ310)"),
    ("Mức phạt cụ thể trong khung (đến hết 15/01/2026)", "Mức trung bình; MỖI tình tiết giảm nhẹ giảm 10% mức trung bình "
     "(không thấp hơn mức tối thiểu); MỖI tình tiết tăng nặng tăng 10% (không vượt mức tối đa); 1 giảm nhẹ trừ 1 tăng nặng.",
     "Điều 7.4.b, 7.4.d NĐ125 gốc"),
    ("Điều 16 - khai sai thiếu thuế", "Phạt 20% số thuế khai thiếu / số thuế miễn, giảm, hoàn cao hơn quy định",
     "Điều 16.1, Điều 7.1.b"),
    ("Điều 17 - trốn thuế (bảng mức)", "≥1 tình tiết giảm nhẹ: 1 lần số thuế trốn | không có tình tiết: 1,5 lần | "
     "1 tình tiết tăng nặng: 2 lần | 2 tình tiết tăng nặng: 2,5 lần | từ 3 tình tiết tăng nặng: 3 lần", "Điều 17.1-5 NĐ125"),
    ("Không xử phạt khi tự khai bổ sung (từ 16/01/2026)", "Khai sai: nếu đã khai bổ sung và nộp đủ thuế TRƯỚC thời điểm cơ quan "
     "thuế công bố quyết định kiểm tra, cơ quan khác công bố quyết định thanh tra/kiểm tra tại trụ sở, hoặc trước khi cơ quan "
     "thuế/cơ quan có thẩm quyền phát hiện thì không bị xử phạt về thuế.", "Điều 9.3 NĐ125 (sửa bởi NĐ310)"),
    ("Gộp hành vi trong cùng một ngày/vụ việc", "Cùng một ngày khai sai nhiều chỉ tiêu: phạt 01 hành vi có khung cao nhất. Cùng "
     "ngày chậm nộp nhiều hồ sơ khai cùng sắc thuế: phạt 01 hành vi cao nhất. Nhiều hóa đơn lập không đúng thời điểm / không "
     "lập trong một vụ việc: phạt 01 hành vi theo tổng số hóa đơn.", "Điều 5.3 NĐ125 (sửa bởi NĐ310)"),
    ("Quy mô lớn (đến hết 15/01/2026)", "Thuế: số thuế (thiếu/trốn/miễn giảm hoàn cao hơn) từ 100.000.000 đ hoặc giá trị hàng "
     "hóa dịch vụ từ 500.000.000 đ trở lên. Hóa đơn: từ 10 số hóa đơn trở lên.", "Điều 6.2 NĐ125 gốc"),
    ("Quy mô lớn (từ 16/01/2026)", "Hóa đơn: từ 10 số hóa đơn trở lên, chỉ đối với hành vi tại khoản 2 Điều 22, Điều 26, Điều 27. "
     "Thuế: chỉ còn trốn thuế từ 100.000.000 đ trở lên khi chuyển hồ sơ xử lý hình sự (tiêu chí 500 triệu giá trị hàng hóa "
     "không còn trong nội dung sửa đổi).", "Điều 1.4 NĐ310 (thay khoản 2 Điều 6)"),
    ("Thời hiệu xử phạt", "Hóa đơn: 01 năm (danh sách hành vi 'đang thực hiện' đã được NĐ310 sửa tại Điều 8.1). Thủ tục thuế: "
     "02 năm. Trốn thuế (chưa đến mức hình sự) và khai sai dẫn đến thiếu thuế: 05 năm.", "Điều 8.1.a NĐ125; Điều 8.2 NĐ125 (sửa bởi NĐ310)"),
    ("Hiệu lực và chuyển tiếp", "NĐ125 hiệu lực 05/12/2020. NĐ310 hiệu lực 16/01/2026. Hành vi đã kết thúc trước 16/01/2026: "
     "áp dụng văn bản có hiệu lực tại thời điểm vi phạm. Hành vi đang thực hiện trước ngày này và bị phát hiện sau ngày này: "
     "áp dụng NĐ310.", "Điều 3 NĐ310"),
    ("Chưa nằm trong bảng này", "Tiền chậm nộp tiền thuế (tỷ lệ theo Luật Quản lý thuế - văn bản chưa được cung cấp, bảng này "
     "không nêu tỷ lệ); truy thu thuế; vi phạm hải quan, bảo hiểm xã hội, lao động; các nghị định xử phạt chuyên ngành khác.", "-"),
]

# ───────────── Xuất file ─────────────
FIELDS = ["lv", "dieu", "khoan", "diem", "hanh_vi", "hinh_thuc", "min_vnd", "max_vnd", "ty_le", "dieu_kien",
          "bien_phap_khac_phuc", "hieu_luc_tu", "hieu_luc_den", "trang_thai", "nguon", "ghi_chu"]
for i, r in enumerate(rows, 1):
    r["id"] = f"P{i:03d}"

(OUT / "penalty_catalog.json").write_text(json.dumps(
    {"rows": rows, "general": [dict(chu_de=a, noi_dung=b, can_cu=c) for a, b, c in general]},
    ensure_ascii=False, indent=1), encoding="utf-8")

with open(OUT / "penalty_catalog.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=["id"] + FIELDS)
    w.writeheader()
    w.writerows([{k: r.get(k, "") for k in ["id"] + FIELDS} for r in rows])

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Checklist mức phạt"
HEAD = ["ID", "Lĩnh vực", "Điều", "Khoản", "Điểm", "Hành vi vi phạm", "Hình thức", "Sàn - tổ chức (đ)", "Trần - tổ chức (đ)",
        "Sàn - cá nhân* (đ)", "Trần - cá nhân* (đ)", "Mức theo tỷ lệ", "Điều kiện / số lượng", "Biện pháp khắc phục hậu quả",
        "Áp dụng từ", "Áp dụng đến", "Trạng thái", "Nguồn", "Ghi chú"]
ws.append(HEAD)
for r in rows:
    mn, mx = r["min_vnd"], r["max_vnd"]
    ws.append([r["id"], r["lv"], r["dieu"], r["khoan"], r["diem"], r["hanh_vi"], r["hinh_thuc"], mn, mx,
               None if mn is None else mn // 2, None if mx is None else mx // 2, r["ty_le"], r["dieu_kien"],
               r["bien_phap_khac_phuc"], r["hieu_luc_tu"], r["hieu_luc_den"], r["trang_thai"], "NĐ " + r["nguon"].replace("125", "125/2020").replace("310", "310/2025"), r["ghi_chu"]])
hdr_fill = PatternFill("solid", fgColor="1F3864")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True, vertical="center")
widths = [7, 14, 6, 6, 6, 60, 15, 15, 15, 15, 15, 30, 40, 36, 12, 12, 28, 16, 50]
for i, w_ in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w_
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for idx in (7, 8, 9, 10):
        row[idx].number_format = "#,##0"
    if row[16].value != "Còn hiệu lực":
        for c in row:
            c.fill = PatternFill("solid", fgColor="F2F2F2")
ws.freeze_panes = "G2"
ws.auto_filter.ref = ws.dimensions

ws2 = wb.create_sheet("Nguyên tắc chung")
ws2.append(["Chủ đề", "Nội dung", "Căn cứ"])
for a, b, c in general:
    ws2.append([a, b, c])
for c in ws2[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = hdr_fill
for col, w_ in zip("ABC", (38, 110, 50)):
    ws2.column_dimensions[col].width = w_
for row in ws2.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")

ws3 = wb.create_sheet("Đọc trước khi dùng")
for line in [
    "Bảng này CHỈ là tra cứu khung phạt theo văn bản (NĐ 125/2020 sửa đổi bởi NĐ 310/2025). Không phải tư vấn pháp lý.",
    "Mức phạt thực tế do cơ quan có thẩm quyền quyết định theo từng hồ sơ; chưa gồm tiền chậm nộp và truy thu thuế.",
    "Cột 'cá nhân*' = 1/2 mức tổ chức (Điều 5.5 NĐ125: tổ chức gấp 2 lần cá nhân, trừ Điều 16, 17, 18); xem sheet Nguyên tắc chung.",
    "Dòng nền xám = quy định đã bị bãi bỏ / thay thế từ 16/01/2026, chỉ áp dụng cho hành vi đã kết thúc trước ngày này.",
    "Nguồn: PDF NĐ125 (Công báo 1017+1018), NĐ310 bản OCR (có lỗi nhận dạng, cần đối chiếu PDF gốc).",
    "Trạng thái kiểm chứng: xem verify_report.txt (kiểm tra cơ học) và báo cáo kiểm chứng độc lập.",
]:
    ws3.append([line])
ws3.column_dimensions["A"].width = 150
wb.move_sheet("Đọc trước khi dùng", offset=-2)
wb.save(OUT / "penalty_catalog.xlsx")
print(f"rows={len(rows)} general={len(general)}")
